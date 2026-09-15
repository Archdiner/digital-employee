"""Plain text out of the files a finance team actually sends. No OCR, no cleverness."""
import io

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        return _docx(data)
    if name.endswith((".xlsx", ".xlsm")):
        return _xlsx(data)
    if name.endswith(".pdf"):
        return _pdf(data)
    if name.endswith(".pptx"):
        return _pptx(data)
    return data.decode("utf-8", errors="replace")


def _docx(data):
    doc = Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for ti, table in enumerate(doc.tables, 1):
        out.append(f"\n[table {ti}]")
        for row in table.rows:
            out.append("\t".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _xlsx(data):
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n[sheet: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if any(v is not None for v in row):
                out.append(f"row {i}\t" + "\t".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _pdf(data):
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(f"\n[page {i}]\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1))


def _pptx(data):
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n[slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append("\t".join(c.text for c in row.cells))
    return "\n".join(out)
