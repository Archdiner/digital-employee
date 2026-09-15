"""End to end against a running server: create employee, upload fixtures, wait for facts, ask for a review,
wait for the document, ask where a number came from.
  Local:    python tests/e2e.py http://localhost:8080
  Deployed: ADMIN_PASSWORD=$(az keyvault secret show --vault-name zybit-emp-kv -n admin-password --query value -o tsv) \
            python tests/e2e.py https://emp-app.jollymushroom-dec64f0f.centralus.azurecontainerapps.io"""
import base64
import os
import re
import sys
import time
import urllib.request
import urllib.parse
import http.cookiejar
import uuid
from pathlib import Path

BASE = sys.argv[1].rstrip("/") if len(sys.argv) > 1 else "http://localhost:8080"
PASSWORD = os.environ.get("ADMIN_PASSWORD")  # set for the deployed app
FIX = Path(__file__).parent / "fixtures"
opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
AUTH = [("Authorization", "Basic " + base64.b64encode(f"admin:{PASSWORD}".encode()).decode())] if PASSWORD else []
opener.addheaders = [("User-Agent", "e2e")] + AUTH


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None


noredir = urllib.request.build_opener(NoRedirect)
noredir.addheaders = [("User-Agent", "e2e")] + AUTH


def get(path):
    return opener.open(BASE + path, timeout=120).read().decode()


def post(path, data=None, files=None):
    if files:
        boundary = uuid.uuid4().hex
        body = b""
        for k, v in (data or {}).items():
            body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode()
        for name, p in files:
            body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"; filename=\"{p.name}\"\r\nContent-Type: application/octet-stream\r\n\r\n".encode() + p.read_bytes() + b"\r\n"
        body += f"--{boundary}--\r\n".encode()
        req = urllib.request.Request(BASE + path, data=body, headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    else:
        req = urllib.request.Request(BASE + path, data=urllib.parse.urlencode(data or {}, doseq=True).encode())
    try:
        r = noredir.open(req, timeout=600)
        return r.status, r.headers.get("Location", ""), r.read().decode()
    except urllib.error.HTTPError as e:
        if e.code in (301, 302, 303, 307):
            return e.code, e.headers.get("Location", ""), ""
        raise


def step(msg):
    print(f"\n== {msg}", flush=True)


step("create employee")
general = re.findall(r'name="skill_ids" value="(\d+)"', get("/employees/new"))
assert general, "no general skills seeded"
_, loc, _ = post("/employees", {
    "name": "Summer Analyst", "firm": "Test Fund Partners", "model": "gpt-5.6-sol",
    "role": "Write the quarterly portfolio review for the CFO",
    "firm_notes": "Fund III has two remaining companies: Northwind Foods and Atlas Logistics.\n"
                  "Quarterly numbers: one management accounts workbook per company, file ending _FINAL, Summary tab, USD 000s. The Commentary tab explains movements.\n"
                  "EBITDA means Adjusted EBITDA as shown in the pack.\n"
                  "The review compares the quarter to the prior quarter and to budget. Match the layout of the last review.",
    "skill_ids": general,
})
eid = int(loc.rsplit("/", 1)[1])
print("employee", eid)

step("upload sources and past review")
post(f"/employees/{eid}/documents", {"kind": "source"}, [("files", FIX / "Northwind_Q2_2026_MA_FINAL.xlsx"), ("files", FIX / "Atlas_Q2_2026_MA_FINAL.xlsx")])
post(f"/employees/{eid}/documents", {"kind": "past_review"}, [("files", FIX / "Fund_III_Q1_2026_Portfolio_Review.docx")])

step("wait for fact extraction")
for _ in range(60):
    html = get(f"/employees/{eid}")
    counts = [int(x) for x in re.findall(r"chars</td><td>(\d+)</td>", html)]
    if len(counts) >= 3 and all(c > 0 for c in counts):
        break
    time.sleep(5)
print("facts per document:", counts)
assert len(counts) >= 3 and all(c > 0 for c in counts), "fact extraction did not finish"

step("ask for the review")
_, loc, _ = post(f"/employees/{eid}/runs", {"task": "Write the Q2 2026 quarterly portfolio review for Fund III."})
rid = int(loc.rsplit("/", 1)[1])
print("run", rid)
for i in range(240):
    html = get(f"/runs/{rid}")
    m = re.search(r'class="pill (\w+)"', html)
    status = m.group(1)
    if status == "waiting":
        q = re.search(r'<pre class="note">(.*?)</pre>', html, re.S).group(1)
        print("it asked:", q[:400])
        post(f"/runs/{rid}/answer", {"answer": "Use the figures in the packs as they are. Compare to prior quarter and to budget. Keep the last review's layout."})
    elif status in ("done", "failed"):
        break
    time.sleep(5)
print("status:", status)
note = re.search(r'Its note to the operator</b><pre class="note">(.*?)</pre>', html, re.S)
print("note:", (note.group(1) if note else "(none)")[:1200])
err = re.search(r"<b>Error</b><pre class=\"note\">(.*?)</pre>", html, re.S)
if err:
    print("error:", err.group(1)[:1500])
assert status == "done", "run did not finish"

step("download document")
did = re.search(r'href="/documents/(\d+)">Download', html).group(1)
data = opener.open(f"{BASE}/documents/{did}").read()
out = Path("/tmp/e2e_review.docx"); out.write_bytes(data)
print("docx bytes", len(data), "->", out)
from docx import Document  # noqa: E402
d = Document(str(out))
print("\n".join(p.text for p in d.paragraphs if p.text.strip())[:2500])
for t in d.tables:
    for r in t.rows:
        print(" | ".join(c.text for c in r.cells))

step("where did this number come from?")
_, _, html = post(f"/runs/{rid}/explain", {"question": "Where did the Q2 2026 revenue for Northwind come from?"})
ans = re.search(r'<pre class="note" style="margin-top:12px">(.*?)</pre>', html, re.S)
print(ans.group(1) if ans else html[-800:])
print("\nE2E OK")
