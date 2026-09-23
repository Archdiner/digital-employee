"""The one guarantee worth a test: no number reaches the document unless it came from the fact store."""
import io

import pytest
from docx import Document

from app import render

FACTS = {
    1: {"id": 1, "value": "47310", "unit": "USD 000s"},
    2: {"id": 2, "value": "16.0%", "unit": "%"},
    3: {"id": 3, "value": "(3125.5)", "unit": "USD 000s"},
    4: {"id": 4, "value": "5.5", "unit": "%"},
    5: {"id": 5, "value": "4.3", "unit": "x"},
    6: {"id": 6, "value": "4000", "unit": None},
}


def test_codes_without_units_keep_their_digits():
    paras, _, _ = build("Reconciled to TB account {{fact:6}}.")
    assert paras[-1] == "Reconciled to TB account 4000."


def build(paragraph, table=None):
    spec = {"title": "Fund III Q2 2026 Portfolio Review", "sections": [{"heading": "Summary", "paragraphs": [paragraph], "table": table}]}
    data, cited = render.build(spec, FACTS.get)
    d = Document(io.BytesIO(data))
    return [p.text for p in d.paragraphs if p.text.strip()], d.tables, cited


def test_values_come_from_store_with_separators_and_units():
    paras, _, cited = build("Revenue was {{fact:1}}, up {{fact:4}}. Loss {{fact:3}}. Leverage {{fact:5}}. Margin {{fact:2}}.")
    assert paras[-1] == "Revenue was 47,310 USD 000s, up 5.5%. Loss (3,125.5) USD 000s. Leverage 4.3x. Margin 16.0%."
    assert cited == [1, 2, 3, 4, 5]


def test_bare_reference_drops_unit_in_table_cells():
    _, tables, _ = build("x", {"columns": ["Company", "Revenue (USD 000s)", "Margin"], "rows": [["Northwind", "{{fact:1:v}}", "{{fact:2}}"]]})
    assert [c.text for c in tables[0].rows[1].cells] == ["Northwind", "47,310", "16.0%"]


def test_labels_are_not_numbers():
    paras, _, _ = build("Q2 2026 versus Q1 and FY2025. H1 2026 as at 30 June 2026. See 2. below; the 3rd item.")
    assert paras


@pytest.mark.parametrize("text", ["FY26 forecast update", "Aug-26 actual vs Aug-26 budget", "Sep-Dec run-rate", "FY 2026/27 plan", "1H26 results", "August 2026 pack"])
def test_period_labels_pass(text):
    assert build(text)[0]


@pytest.mark.parametrize("text", ["Revenue was 47,310.", "up 8% on the quarter", "Headcount 431", "Fact {{fact:99}} missing"])
def test_stray_numbers_and_missing_facts_are_rejected(text):
    with pytest.raises(render.ReviewError):
        build(text)
