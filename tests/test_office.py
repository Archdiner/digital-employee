"""pptx and xlsx builders obey the same rule as docx: numbers only through fact references."""
import io

import pytest
from openpyxl import load_workbook
from pptx import Presentation

from app import office, render

FACTS = {1: {"id": 1, "value": "47310", "unit": "USD 000s"}, 2: {"id": 2, "value": "16.0%", "unit": "%"}}


def test_pptx_builds_with_facts():
    spec = {"title": "Fund III Q2 2026", "slides": [
        {"title": "Northwind", "bullets": ["Revenue {{fact:1}}", "Margin {{fact:2}}"]},
        {"title": "Table", "table": {"columns": ["Company", "Revenue (USD 000s)"], "rows": [["Northwind", "{{fact:1:v}}"]]}, "bullets": ["Q2 2026"]},
    ]}
    data, cited = office.build_pptx(spec, FACTS.get)
    prs = Presentation(io.BytesIO(data))
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert any("47,310 USD 000s" in t for t in texts) and cited == [1, 2]
    assert len(prs.slides) == 3


def test_xlsx_builds_numeric_cells():
    spec = {"title": "x", "sheets": [{"name": "Summary", "rows": [["Metric", "Q2 2026"], ["Revenue", "{{fact:1:v}}"], ["Margin", "{{fact:2}}"]]}]}
    data, cited = office.build_xlsx(spec, FACTS.get)
    ws = load_workbook(io.BytesIO(data))["Summary"]
    assert ws["B2"].value == 47310 and ws["B3"].value == "16.0%" and cited == [1, 2]


def test_grid_and_text_substitution_reject_stray_numbers():
    with pytest.raises(render.ReviewError):
        office.substitute_grid([["Revenue", "47310"]], FACTS.get)
    with pytest.raises(render.ReviewError):
        office.substitute_text("grew 8%", FACTS.get)
    assert office.substitute_grid([["Revenue", "{{fact:1:v}}"]], FACTS.get)[0] == [["Revenue", 47310]]
