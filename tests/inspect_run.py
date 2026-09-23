"""Print a run: transcript, files (with contents), rejections, tool sequence. python tests/inspect_run.py <run_id>"""
import io
import json
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app import agent, db

rid = int(sys.argv[1])
r = db.q("select status, error from runs where id = %s", (rid,), one=True)
print(f"=== run {rid}: {r['status']}", ("ERROR: " + r["error"][-600:]) if r["error"] else "")
for m in agent.transcript(rid):
    if m["who"] == "file":
        print(f"[FILE] {m['doc']['filename']}")
    else:
        print(f"[{m['who'].upper()}] {m['text'][:2500]}\n")
acts = [w["action"] for w in db.q("select action from worklog where run_id = %s order by id", (rid,))]
print("tools:", " ".join(acts))
for w in db.q("select detail from worklog where run_id = %s and action like %s order by id", (rid, "%rejected%")):
    print("REJECTED:", str(w["detail"].get("problems"))[:600])
for d in db.q("select filename, bytes from documents where run_id = %s order by id", (rid,)):
    name, data = d["filename"], d["bytes"]
    print(f"\n--- {name} ---")
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        print("\n".join(p.text for p in doc.paragraphs if p.text.strip())[:6000])
        for t in doc.tables:
            for row in t.rows:
                print(" | ".join(c.text for c in row.cells))
    elif name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data))
        for ws in wb.worksheets:
            print(f"[sheet {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                print("  ", row)
    elif name.endswith(".pptx"):
        for i, s in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            print(f"  {i}. " + " / ".join(sh.text_frame.text.replace(chr(10), ' ')[:200] for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()))
